from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colour palette
DARK_NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ACCENT_BLUE = RGBColor(0x2E, 0x86, 0xAB)
ACCENT_TEAL = RGBColor(0x00, 0xB4, 0xD8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xF0, 0xF0, 0xF0)
DARK_GREY = RGBColor(0x33, 0x33, 0x33)
AMBER = RGBColor(0xF4, 0xA2, 0x61)
RED_ACCENT = RGBColor(0xE7, 0x6F, 0x51)
GREEN_ACCENT = RGBColor(0x2A, 0x9D, 0x8F)

def add_bg(slide, color=DARK_NAVY):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text_box(slide, left, top, width, height, text, font_size=18, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_bullet_slide(slide, left, top, width, height, items, font_size=16, color=WHITE):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = Pt(8)
    return txBox

def add_rect(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape

# ============================================================
# SECTION 1: 4-MINUTE PRESENTATION (Slides 1-7)
# ============================================================

# ============================================================
# SLIDE 1 — Title
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, 0, 0, 13.333, 0.08, ACCENT_TEAL)
add_text_box(slide, 1, 1.5, 11, 1.5, "Dispatcher Decision Support Agent", 42, WHITE, True)
add_text_box(slide, 1, 3.2, 11, 1, "Agentic Transformation of Customer Operations", 24, ACCENT_TEAL, False)
add_text_box(slide, 1, 4.5, 11, 0.8, "Apex Distribution Ltd  |  35-person Customer Ops  |  4 Work Streams Assessed", 16, RGBColor(0xAA, 0xAA, 0xAA))
add_text_box(slide, 1, 5.5, 11, 0.8, "Prepared for: Sarah Whitmore, COO  \u2022  Gate 2  \u2022  May 2026", 14, RGBColor(0x88, 0x88, 0x88))

# ============================================================
# SLIDE 2 — The 4 Work Streams: Summary & Why We Chose Exceptions
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, 0, 0, 13.333, 0.08, ACCENT_TEAL)
add_text_box(slide, 0.8, 0.4, 11, 0.7, "CUSTOMER OPERATIONS \u2014 4 WORK STREAMS", 14, ACCENT_TEAL, True)
add_text_box(slide, 0.8, 1.0, 11, 0.8, "All assessed. One wins on Volume \u00d7 Value. Here's why.", 28, WHITE, True)

# Four work stream cards
ws_data = [
    ("Delivery\nExceptions", "180/day \u2022 12 min", "Score: 16", "PRIMARY TARGET", GREEN_ACCENT,
     "High judgment, time-critical,\nCRM + Driver App APIs available,\nno Aurum dependency"),
    ("ETA\nInquiries", "400/day \u2022 4 min", "Score: 10", "AUTOMATION / RPA", RGBColor(0x88, 0x88, 0x88),
     "Mostly deterministic lookup.\nNot agentic \u2014 standard RPA.\nEdge cases only need human."),
    ("Dispatch\nAdjustments", "90/day \u2022 18 min", "Score: 12", "WAVE 2 (API blocked)", AMBER,
     "High value but dispatch console\nhas limited API. Blocked until\nintegration path confirmed."),
    ("Billing\nDisputes", "60/day \u2022 28 min", "Score: 10", "AGENT-ASSIST ONLY", RGBColor(0x88, 0x88, 0x88),
     "Aurum = batch-only, schema breaks.\nRPA already failed here. Credit\ndecisions need policy first."),
]

for i, (title, volume, score, label, accent, reason) in enumerate(ws_data):
    x = 0.8 + i * 3.1
    add_rect(slide, x, 2.0, 2.8, 5.0, RGBColor(0x22, 0x35, 0x55))
    add_rect(slide, x, 2.0, 2.8, 0.06, accent)
    add_text_box(slide, x + 0.2, 2.15, 2.4, 1.0, title, 16, WHITE, True)
    add_text_box(slide, x + 0.2, 3.1, 2.4, 0.5, volume, 12, RGBColor(0xAA, 0xAA, 0xAA))
    add_text_box(slide, x + 0.2, 3.6, 2.4, 0.5, score, 22, accent, True)
    add_text_box(slide, x + 0.2, 4.3, 2.4, 0.5, label, 11, accent, True)
    add_text_box(slide, x + 0.2, 4.9, 2.4, 1.8, reason, 11, RGBColor(0xBB, 0xBB, 0xBB))

add_text_box(slide, 0.8, 7.1, 12, 0.3,
    "Score = Execution Frequency (1\u20135) \u00d7 Non-Deterministic Decision Effort (1\u20135). ETA is high-volume but low-judgment. Billing is high-effort but Aurum-blocked.",
    12, RGBColor(0x88, 0x88, 0x88))

# ============================================================
# SLIDE 3 — BEFORE: Delivery Exceptions (the lived process)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, 0, 0, 13.333, 0.08, RED_ACCENT)
add_text_box(slide, 0.8, 0.4, 11, 0.7, "BEFORE \u2014 DELIVERY EXCEPTIONS TODAY", 14, RED_ACCENT, True)
add_text_box(slide, 0.8, 1.0, 11, 0.8, "180/day. 12 min/case. Dispatchers spend 5 min on lookup before they can think.", 26, WHITE, True)

# Left column — What dispatchers do
add_rect(slide, 0.8, 2.2, 5.8, 4.8, RGBColor(0x22, 0x35, 0x55))
add_rect(slide, 0.8, 2.2, 5.8, 0.06, RED_ACCENT)
add_text_box(slide, 1.0, 2.35, 5.4, 0.4, "What dispatchers do (lived, not SOP)", 15, RED_ACCENT, True)
add_bullet_slide(slide, 1.0, 2.85, 5.4, 4.0, [
    "\u2022 Driver calls/voicemails dispatch with exception",
    "\u2022 Dispatcher opens 3 systems: CRM, dispatch console, GPS",
    "\u2022 Pulls route, account, history \u2014 ~5 min context assembly",
    "\u2022 Triage decision under time pressure (driver parked)",
    "\u2022 Verbal instruction back to driver",
    "\u2022 Exception logged in console (often delayed/forgotten)",
    "\u2022 No pattern visibility \u2014 each case isolated",
    "\u2022 SOP is stale: DispatchHub retired Oct 2024; \u00a74.3 = 'TBD'",
], 13, RGBColor(0xDD, 0xDD, 0xDD))

# Right column — Bottlenecks and failures
add_rect(slide, 6.9, 2.2, 5.8, 4.8, RGBColor(0x22, 0x35, 0x55))
add_rect(slide, 6.9, 2.2, 5.8, 0.06, AMBER)
add_text_box(slide, 7.1, 2.35, 5.4, 0.4, "Bottlenecks & failure points", 15, AMBER, True)
add_bullet_slide(slide, 7.1, 2.85, 5.4, 4.0, [
    "\u2022 36 person-hours/day on exceptions alone",
    "\u2022 ~15 hrs/day on context assembly (not judgment)",
    "\u2022 Mark parked 15+ min \u2014 Sandra's line busy, 6 drops late",
    "\u2022 Account sensitivity lives in heads ('the big one')",
    "\u2022 No damage protocol \u2014 \u00a74.3 empty since 2023",
    "\u2022 Dispatch console: Citrix/Java, limited API",
    "\u2022 Driver-vs-customer conflict with zero visual evidence",
    "\u2022 No escalation protocol beyond stale \u00a3500 threshold",
], 13, RGBColor(0xDD, 0xDD, 0xDD))

# ============================================================
# SLIDE 4 — AFTER: The Redesigned Process with DDSA
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, 0, 0, 13.333, 0.08, GREEN_ACCENT)
add_text_box(slide, 0.8, 0.4, 11, 0.7, "AFTER \u2014 WITH THE DDSA IN PLACE", 14, GREEN_ACCENT, True)
add_text_box(slide, 0.8, 1.0, 11, 0.8, "Agent assembles context in <30s \u2014 dispatchers focus on judgment", 26, WHITE, True)

# Three-column layout
arch_data = [
    ("AGENT DECIDES ALONE\n(context assembly)", GREEN_ACCENT, [
        "Detect inbound exception signal",
        "Pull account record from CRM (API)",
        "Pull route data + GPS position",
        "Calculate route pressure score",
        "Classify account sensitivity (rule-based)",
        "Assemble context card (<30s SLA)",
        "Queue priority if driver wait >15 min",
    ]),
    ("HUMAN-LED + AGENT SUPPORT\n(triage decision)", ACCENT_TEAL, [
        "Present triage options with trade-offs",
        "Surface 'repeat pattern' alert",
        "Propose Duty Manager escalation (>\u00a3500)",
        "Flag unauthorised refusing party",
        "Dispatcher decides: return/leave/hold",
        "Dispatcher instructs driver",
        "Dispatcher logs and escalates",
    ]),
    ("HUMAN-ONLY\n(cannot be delegated)", AMBER, [
        "Assess damage credibility",
        "  \u2192 verbal-only, no photos, liability",
        "Decide credit amounts",
        "  \u2192 no policy exists; pure discretion",
        "",
        "WHY: Input Structure = L",
        "Decision Determinism = L",
        "Risk/Compliance = H",
    ]),
]

for i, (title, accent, items) in enumerate(arch_data):
    x = 0.8 + i * 4.1
    add_rect(slide, x, 2.1, 3.8, 5.0, RGBColor(0x22, 0x35, 0x55))
    add_rect(slide, x, 2.1, 3.8, 0.06, accent)
    add_text_box(slide, x + 0.2, 2.25, 3.4, 0.7, title, 12, accent, True)
    add_bullet_slide(slide, x + 0.2, 3.0, 3.4, 4.0, ["\u2022 " + item if item else "" for item in items], 12, RGBColor(0xDD, 0xDD, 0xDD))

add_text_box(slide, 0.8, 7.2, 12, 0.3,
    "Zero tasks are fully agentic end-to-end. This is a decision-support agent \u2014 not a chatbot, not RPA. Makes dispatchers faster, not redundant.",
    12, RGBColor(0x88, 0x88, 0x88))

# ============================================================
# SLIDE 5 — The Agent: DDSA Identity & KPIs
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, 0, 0, 13.333, 0.08, ACCENT_TEAL)
add_text_box(slide, 0.8, 0.4, 11, 0.7, "THE AGENT \u2014 DDSA", 14, ACCENT_TEAL, True)
add_text_box(slide, 0.8, 1.0, 11, 0.8, "Dispatcher Decision Support Agent", 32, WHITE, True)

add_text_box(slide, 0.8, 2.0, 11, 0.8,
    "Job: Assemble decision context (account, route, history, value) in seconds not minutes,\n"
    "so dispatchers make better triage calls faster. The agent does NOT make triage decisions.", 15, RGBColor(0xDD, 0xDD, 0xDD))

# KPIs
add_rect(slide, 0.8, 3.0, 12.0, 0.06, ACCENT_TEAL)
add_text_box(slide, 0.8, 3.15, 12, 0.4, "KPI TARGETS", 13, ACCENT_TEAL, True)

kpi_data = [
    ("<30s", "Context\nAssembly", "Signal to card"),
    ("90%", "Context\nCompleteness", "No manual lookup"),
    ("80%", "Dispatcher\nAdoption", "Within 4 weeks"),
    ("25%", "Handling Time\nReduction", "12 min \u2192 9 min"),
    ("<10%", "False-Alert\nRate", "Flags overridden"),
]

for i, (value, label, desc) in enumerate(kpi_data):
    x = 0.8 + i * 2.4
    add_text_box(slide, x, 3.5, 2.2, 0.6, value, 28, ACCENT_TEAL, True, PP_ALIGN.CENTER)
    add_text_box(slide, x, 4.2, 2.2, 0.6, label, 12, WHITE, True, PP_ALIGN.CENTER)
    add_text_box(slide, x, 4.8, 2.2, 0.5, desc, 11, RGBColor(0xAA, 0xAA, 0xAA), False, PP_ALIGN.CENTER)

# Why this avoids prior failures
add_rect(slide, 0.8, 5.5, 12.0, 0.06, ACCENT_TEAL)
add_text_box(slide, 0.8, 5.65, 12, 0.4, "WHY THIS AVOIDS PRIOR FAILURE MODES", 13, ACCENT_TEAL, True)
add_bullet_slide(slide, 0.8, 6.0, 12, 1.5, [
    "1. Not a chatbot \u2014 internal tool for dispatchers, not customer-facing. 2024 chatbot failure won't repeat.",
    "2. Doesn't touch Aurum \u2014 avoids schema fragility that killed the RPA project. Zero legacy dependency.",
    "3. Human stays in the loop \u2014 agent supports, never decides. Sarah's team keeps authority, gains speed.",
], 13, RGBColor(0xDD, 0xDD, 0xDD))

# ============================================================
# SLIDE 6 — Defending Delegation Calls
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, 0, 0, 13.333, 0.08, ACCENT_TEAL)
add_text_box(slide, 0.8, 0.4, 11, 0.7, "DEFENDING THE DELEGATION CALLS", 14, ACCENT_TEAL, True)
add_text_box(slide, 0.8, 1.0, 11, 0.8, "Why context assembly is safe to fully delegate \u2014 and why triage is not", 26, WHITE, True)

defences = [
    ("Why is context assembly safe for full delegation?",
     "Read-only lookup across CRM + Driver App + GPS. No judgment, no customer contact, no financial impact.\n"
     "Input Structure = H. Decision Determinism = H. Risk = L. Reversible (card can refresh)."),
    ("Why isn't the triage decision itself agentic?",
     "Input Structure = L (verbal-only driver signal). Decision Determinism = L (every refusal is different).\n"
     "Risk = H (wrong call = lost \u00a32K pallet or lost key account). No visual evidence. Liability sits with human."),
    ("Why is credit amount human-only, not even agent-assisted?",
     "No documented credit policy exists. Sandra chose \u00a3170/\u00a3340 (50%) with no formula. Another credit was \u00a388 \u2014 also no rule.\n"
     "Until a policy exists, any agent suggestion creates compliance risk. Adjacent archetype rejected."),
    ("What about the dispatch console's limited API?",
     "Primary integration risk. If route data inaccessible \u2192 context card shows 'Route: unavailable'.\n"
     "Agent degrades gracefully \u2014 never blocks dispatcher. Partial card > no card."),
]

for i, (question, answer) in enumerate(defences):
    y = 2.0 + i * 1.3
    add_rect(slide, 0.8, y, 11.7, 1.15, RGBColor(0x22, 0x35, 0x55))
    add_text_box(slide, 1.0, y + 0.05, 11.3, 0.4, question, 14, AMBER, True)
    add_text_box(slide, 1.0, y + 0.45, 11.3, 0.7, answer, 12, RGBColor(0xDD, 0xDD, 0xDD))

# ============================================================
# SLIDE 7 — Billing Disputes: Why NOT primary + Wave 2
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, 0, 0, 13.333, 0.08, AMBER)
add_text_box(slide, 0.8, 0.4, 11, 0.7, "BILLING DISPUTES \u2014 WHY NOT PRIMARY TARGET", 14, AMBER, True)
add_text_box(slide, 0.8, 1.0, 11, 0.8, "Mapped in D1, scored in D2 \u2014 deliberately NOT chosen for Wave 1", 26, WHITE, True)

# Left — what we found
add_rect(slide, 0.8, 2.2, 5.8, 4.8, RGBColor(0x22, 0x35, 0x55))
add_rect(slide, 0.8, 2.2, 5.8, 0.06, RED_ACCENT)
add_text_box(slide, 1.0, 2.35, 5.4, 0.4, "What the cognitive map revealed", 15, RED_ACCENT, True)
add_bullet_slide(slide, 1.0, 2.85, 5.4, 4.0, [
    "\u2022 60/day, 28 min/case = 28 person-hours/day",
    "\u2022 9-day resolution cycle (Artefact 2)",
    "\u2022 Billing deflects to Ops; customer re-explains",
    "\u2022 Sandra applies credits via manual override",
    "\u2022 No audit trail (untracked \u00a3170 goodwill credit)",
    "\u2022 No credit policy \u2014 amount is pure discretion",
    "\u2022 Cross-references delivery exceptions \u2194 invoices",
    "\u2022 Aurum: batch-only, 24h lag, schema changes quarterly",
], 13, RGBColor(0xDD, 0xDD, 0xDD))

# Right — why not and what future looks like
add_rect(slide, 6.9, 2.2, 5.8, 4.8, RGBColor(0x22, 0x35, 0x55))
add_rect(slide, 6.9, 2.2, 5.8, 0.06, AMBER)
add_text_box(slide, 7.1, 2.35, 5.4, 0.4, "Why it loses + Wave 2 plan", 15, AMBER, True)
add_bullet_slide(slide, 7.1, 2.85, 5.4, 4.0, [
    "\u2022 Aurum has NO API \u2014 RPA already failed here",
    "\u2022 Schema changes quarterly without notice",
    "\u2022 Credit decisions = Human-only (no policy)",
    "\u2022 V\u00d7V score: 10 vs 16 for Exceptions",
    "",
    "\u2022 Wave 2 agent-assist opportunity:",
    "  \u2192 Dispute type classification (Agent-led)",
    "  \u2192 Invoice-to-exception cross-reference",
    "  \u2192 Customer notification drafting",
    "  \u2192 Requires: credit policy + Aurum API/wrapper",
], 13, RGBColor(0xDD, 0xDD, 0xDD))

# ============================================================
# SECTION 2: 3-MINUTE COO QUESTION (Slides 8-9)
# ============================================================

# ============================================================
# SLIDE 8 — The One Question for the COO
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, 0, 0, 13.333, 0.08, AMBER)
add_text_box(slide, 0.8, 0.4, 11, 0.7, "COO QUESTION \u2014 3 MINUTES", 14, AMBER, True)
add_text_box(slide, 0.8, 1.0, 11, 1.0, "The question that would most change my design", 32, WHITE, True)

add_rect(slide, 0.8, 2.3, 11.7, 1.5, RGBColor(0x22, 0x35, 0x55))
add_rect(slide, 0.8, 2.3, 11.7, 0.06, AMBER)
add_text_box(slide, 1.0, 2.5, 11.3, 0.4, "THE QUESTION", 14, AMBER, True)
add_text_box(slide, 1.0, 3.0, 11.3, 0.8,
    "\"When Mark called in about the Stein-Allen pallet \u2014 how long did it take Sandra to pull up his route, "
    "the account, and decide what to tell him? Roughly?\"", 18, WHITE, True)

# Why this question
add_rect(slide, 0.8, 4.1, 5.6, 3.0, RGBColor(0x22, 0x35, 0x55))
add_text_box(slide, 1.0, 4.2, 5.2, 0.4, "Why this question matters", 14, GREEN_ACCENT, True)
add_bullet_slide(slide, 1.0, 4.7, 5.2, 2.3, [
    "\u2022 Entire DDSA value prop rests on context assembly",
    "  taking ~5 min of 12 min/case [ASSUMED \u2014 A4]",
    "\u2022 If actually 1 min (Sandra knows routes cold),",
    "  agent saves negligible time",
    "\u2022 KPI target (25% reduction) only works if",
    "  context assembly IS the bottleneck",
], 13, RGBColor(0xDD, 0xDD, 0xDD))

# What changes based on answer
add_rect(slide, 6.7, 4.1, 5.8, 3.0, RGBColor(0x22, 0x35, 0x55))
add_text_box(slide, 6.9, 4.2, 5.4, 0.4, "What changes based on the answer", 14, ACCENT_TEAL, True)
add_bullet_slide(slide, 6.9, 4.7, 5.4, 2.3, [
    "If >5 min \u2192 design validated, proceed",
    "If 2\u20135 min \u2192 still viable, adjust KPI targets down",
    "If <2 min \u2192 pivot value prop from speed to",
    "  pattern detection (repeat accounts, trends)",
    "If 'it varies' \u2192 push for new-dispatcher vs",
    "  experienced-dispatcher split",
], 13, RGBColor(0xDD, 0xDD, 0xDD))

# ============================================================
# SLIDE 9 — Follow-Up Probes for the COO
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, 0, 0, 13.333, 0.08, AMBER)
add_text_box(slide, 0.8, 0.4, 11, 0.7, "COO QUESTION \u2014 FOLLOW-UP PROBES", 14, AMBER, True)
add_text_box(slide, 0.8, 1.0, 11, 0.8, "How to push if Sarah hedges or deflects", 28, WHITE, True)

probes = [
    ("If she says 'it varies':",
     "\"Two cases \u2014 fastest and slowest. For a dispatcher here 6 months vs Sandra who knows everything \u2014 how different?\"",
     "New-dispatcher time validates agent value; experienced time shows the ceiling"),
    ("If she says 'Sandra just knows':",
     "\"What about when Sandra's not there? Mark couldn't reach her. What happens to the others?\"",
     "Tests whether problem is team-wide or Sandra-specific"),
    ("If she gives a number:",
     "\"How much is looking things up vs actually thinking? I want the lookup time separately.\"",
     "Separates context assembly (agent can help) from judgment (agent cannot)"),
    ("If she pushes back ('why does that matter?'):",
     "\"My design is a context tool, not a decision tool. If lookup isn't the bottleneck, I'm solving the wrong problem.\"",
     "Shows honesty \u2014 builds trust by admitting design could be wrong"),
]

for i, (scenario, probe, rationale) in enumerate(probes):
    y = 2.0 + i * 1.3
    add_rect(slide, 0.8, y, 11.7, 1.15, RGBColor(0x22, 0x35, 0x55))
    add_text_box(slide, 1.0, y + 0.05, 11.3, 0.35, scenario, 13, AMBER, True)
    add_text_box(slide, 1.0, y + 0.4, 11.3, 0.4, probe, 12, WHITE)
    add_text_box(slide, 1.0, y + 0.85, 11.3, 0.3, rationale, 11, RGBColor(0xAA, 0xAA, 0xAA))

# ============================================================
# SECTION 3: 3-MINUTE COACH Q&A (Slides 10-12)
# ============================================================

# ============================================================
# SLIDE 10 — Coach Q&A: Key Discovery Questions
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, 0, 0, 13.333, 0.08, ACCENT_BLUE)
add_text_box(slide, 0.8, 0.4, 11, 0.7, "COACH Q&A \u2014 TOP DISCOVERY QUESTIONS", 14, ACCENT_BLUE, True)
add_text_box(slide, 0.8, 1.0, 11, 0.8, "Questions whose answers would materially change the design", 28, WHITE, True)

questions = [
    ("Q1", "Context assembly bottleneck", "How long does context assembly actually take? (A4 \u2014 entire value prop depends on this)", GREEN_ACCENT),
    ("Q2", "Driver App API capability", "Can the Driver App push signals outward? (S1 \u2014 real-time triggering depends on this)", GREEN_ACCENT),
    ("Q3", "Dispatch console API surface", "What can we read from the Java/Citrix console? (S2 \u2014 route data accessibility)", ACCENT_TEAL),
    ("Q4", "\u00a3500 escalation threshold", "Is the SOP threshold still real? (A7 \u2014 ET-1 built around this number)", ACCENT_TEAL),
    ("Q5", "What killed prior AI projects", "Chatbot + RPA \u2014 technology failure or team rejection? (A9 \u2014 adoption risk)", AMBER),
]

y_start = 2.2
for i, (code, topic, detail, accent) in enumerate(questions):
    y = y_start + i * 1.0
    bg_color = RGBColor(0x22, 0x35, 0x55) if i % 2 == 0 else RGBColor(0x1E, 0x30, 0x50)
    add_rect(slide, 0.8, y, 11.7, 0.85, bg_color)
    add_text_box(slide, 1.0, y + 0.1, 0.8, 0.4, code, 14, accent, True)
    add_text_box(slide, 2.0, y + 0.1, 3.5, 0.4, topic, 14, WHITE, True)
    add_text_box(slide, 2.0, y + 0.45, 10.0, 0.4, detail, 12, RGBColor(0xAA, 0xAA, 0xAA))

add_text_box(slide, 0.8, 7.2, 12, 0.3, "Priority 1 (Q1\u2013Q2): design-breaking if wrong.  Priority 2 (Q3\u2013Q5): design-changing but recoverable.", 13, RGBColor(0x88, 0x88, 0x88))

# ============================================================
# SLIDE 11 — Coach Q&A: Assumptions to Defend
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, 0, 0, 13.333, 0.08, ACCENT_BLUE)
add_text_box(slide, 0.8, 0.4, 11, 0.7, "ASSUMPTIONS TO DEFEND", 14, ACCENT_BLUE, True)
add_text_box(slide, 0.8, 1.0, 11, 0.8, "Key assumptions \u2014 confidence levels and what breaks if wrong", 28, WHITE, True)

assumptions = [
    ("A4", "Medium", "Context assembly takes ~5 min of 12 min/case", "Agent value shifts from speed \u2192 pattern detection"),
    ("S1", "Medium", "Driver App has outbound API/webhook", "Need polling fallback (+2\u20135 min latency)"),
    ("S2", "Low", "Dispatch console has read API for route data", "Route pressure score unavailable; partial card only"),
    ("A6", "Medium", "Account sensitivity is informal (not in system)", "If it IS flagged: agent classification is redundant"),
    ("A9", "Medium", "Sarah's scepticism = customer-facing AI only", "If extends to internal: adoption risk is existential"),
]

# Header
add_rect(slide, 0.8, 2.2, 11.7, 0.5, RGBColor(0x15, 0x22, 0x3D))
add_text_box(slide, 1.0, 2.25, 0.8, 0.4, "ID", 12, ACCENT_TEAL, True)
add_text_box(slide, 1.9, 2.25, 1.2, 0.4, "Confidence", 12, ACCENT_TEAL, True)
add_text_box(slide, 3.3, 2.25, 4.5, 0.4, "Assumption", 12, ACCENT_TEAL, True)
add_text_box(slide, 8.0, 2.25, 4.3, 0.4, "If wrong, what breaks", 12, ACCENT_TEAL, True)

for i, (aid, conf, assumption, breaks) in enumerate(assumptions):
    y = 2.8 + i * 0.85
    bg_color = RGBColor(0x22, 0x35, 0x55) if i % 2 == 0 else RGBColor(0x1E, 0x30, 0x50)
    conf_color = AMBER if conf == "Medium" else RED_ACCENT
    add_rect(slide, 0.8, y, 11.7, 0.75, bg_color)
    add_text_box(slide, 1.0, y + 0.15, 0.8, 0.4, aid, 13, WHITE, True)
    add_text_box(slide, 1.9, y + 0.15, 1.2, 0.4, conf, 12, conf_color, True)
    add_text_box(slide, 3.3, y + 0.15, 4.5, 0.5, assumption, 12, RGBColor(0xDD, 0xDD, 0xDD))
    add_text_box(slide, 8.0, y + 0.15, 4.3, 0.5, breaks, 12, RGBColor(0xAA, 0xAA, 0xAA))

add_text_box(slide, 0.8, 7.1, 12, 0.3, "No assumption is 'High' confidence \u2014 all require validation. High = coach-session-validated only.", 13, RGBColor(0x88, 0x88, 0x88))

# ============================================================
# SLIDE 12 — Coach Q&A: Honest Gaps & Trade-offs
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, 0, 0, 13.333, 0.08, ACCENT_BLUE)
add_text_box(slide, 0.8, 0.4, 11, 0.7, "HONEST GAPS & TRADE-OFFS", 14, ACCENT_BLUE, True)
add_text_box(slide, 0.8, 1.0, 11, 0.8, "What I don't know \u2014 and what I'd change with more information", 26, WHITE, True)

# Known gaps
add_rect(slide, 0.8, 2.2, 5.8, 2.3, RGBColor(0x22, 0x35, 0x55))
add_rect(slide, 0.8, 2.2, 5.8, 0.06, RED_ACCENT)
add_text_box(slide, 1.0, 2.35, 5.4, 0.4, "KNOWN GAPS", 14, RED_ACCENT, True)
add_bullet_slide(slide, 1.0, 2.8, 5.4, 1.8, [
    "\u2022 Only mapped 2 of 4 work streams in depth",
    "\u2022 Consignment value source unknown",
    "\u2022 Voicemail vs app channel split not confirmed",
    "\u2022 No data on exception type breakdown",
    "\u2022 Dispatch console API surface unverified",
], 12, RGBColor(0xDD, 0xDD, 0xDD))

# Trade-offs
add_rect(slide, 6.9, 2.2, 5.8, 2.3, RGBColor(0x22, 0x35, 0x55))
add_rect(slide, 6.9, 2.2, 5.8, 0.06, AMBER)
add_text_box(slide, 7.1, 2.35, 5.4, 0.4, "DELIBERATE TRADE-OFFS", 14, AMBER, True)
add_bullet_slide(slide, 7.1, 2.8, 5.4, 1.8, [
    "\u2022 Support-agent over autonomous: politically safer",
    "\u2022 Avoided Aurum entirely: RPA failure lesson",
    "\u2022 Conservative sensitivity rule: fewer alerts > noise",
    "\u2022 Wave 1 avoids deep console integration",
    "\u2022 Billing disputes deferred to Wave 2",
], 12, RGBColor(0xDD, 0xDD, 0xDD))

# What changes
add_rect(slide, 0.8, 4.8, 11.7, 2.4, RGBColor(0x22, 0x35, 0x55))
add_rect(slide, 0.8, 4.8, 11.7, 0.06, GREEN_ACCENT)
add_text_box(slide, 1.0, 4.95, 11.3, 0.4, "IF SARAH'S ANSWERS CHANGE MY DESIGN", 14, GREEN_ACCENT, True)
add_bullet_slide(slide, 1.0, 5.4, 11.3, 1.7, [
    "\u2022 Context assembly <2 min \u2192 pivot to pattern detection agent (repeat accounts, escalation prediction)",
    "\u2022 Driver App has no API \u2192 voicemail transcription as primary signal; accept 2\u20135 min latency; redesign SLA",
    "\u2022 Dispatch console fully locked \u2192 remove route pressure from card; derive rough data from GPS only",
    "\u2022 Sarah's scepticism extends to internal tools \u2192 need dispatcher champion program BEFORE building",
], 12, RGBColor(0xDD, 0xDD, 0xDD))

# ============================================================
# Save
# ============================================================
output_path = r"C:\Users\ChrisEdwards1\IdeaProjects\FDE_Training\Week2\Gate2\Gate2-Presentation-Chris-Edwards-v2.pptx"
prs.save(output_path)
print(f"Saved to {output_path}")
